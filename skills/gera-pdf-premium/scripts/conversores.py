# -*- coding: utf-8 -*-
"""
gera-pdf-premium · conversores: qualquer formato → Markdown normalizado (fonte.md)

Formatos: .md/.markdown (cópia byte a byte), .txt, .docx, .pdf com texto e .pdf
escaneado (OCR com tesseract). Toda limpeza técnica de extração — cabeçalhos e
rodapés corridos, fólios, sumário antigo, hifenização de fim de linha, linhas
duplicadas — é CONVERSÃO, não edição de conteúdo, e fica declarada em
conversao.md com contagem e exemplos.
"""

import collections
import html
import re
import shutil
import statistics
import subprocess
import tempfile
import unicodedata
from pathlib import Path

FORMATOS = {".md": "markdown", ".markdown": "markdown", ".txt": "texto", ".text": "texto",
            ".docx": "docx", ".pdf": "pdf", ".srt": "legenda", ".vtt": "legenda"}


class ErroConversao(Exception):
    pass


# ============================================================
# ESCAPE — texto puro vindo de TXT/DOCX/PDF não pode virar marcação
# ============================================================
def esc_inline(s):
    s = s.replace("\\", "\\\\")
    s = re.sub(r"([*_`\[\]<>])", r"\\\1", s)
    return re.sub(r"&(?=#?\w+;)", r"\\&", s)


def esc_bloco(linha):
    """Escapa o que, no começo da linha, o Markdown leria como estrutura."""
    m = re.match(r"^(\s*)(#{1,6}(?=\s|$)|>|[-+*](?=\s)|\d{1,9}(?=[.)]\s)|={3,}\s*$|-{3,}\s*$|\|)", linha)
    if not m:
        return linha
    ini, marca = m.group(1), m.group(2)
    if marca[0].isdigit():  # "1. texto" → "1\. texto"
        return ini + marca + "\\" + linha[len(ini) + len(marca):]
    return ini + "\\" + linha[len(ini):]


def esc_texto(s):
    return "\n".join(esc_bloco(esc_inline(l)) for l in s.split("\n"))


# ============================================================
# REGISTRO DA LIMPEZA DECLARADA
# ============================================================
class Registro:
    def __init__(self):
        self.itens = collections.OrderedDict()
        self.avisos = []

    def add(self, acao, exemplo=None, n=1):
        r = self.itens.setdefault(acao, {"acao": acao, "n": 0, "exemplos": []})
        r["n"] += n
        if exemplo and len(r["exemplos"]) < 4 and exemplo not in r["exemplos"]:
            r["exemplos"].append(exemplo[:120])

    def lista(self):
        return list(self.itens.values())


def montar_md(blocos):
    """blocos: dicts {tipo: h|p|li|ol|img|tabela|hr, texto, nivel}. Listas contíguas ficam juntas."""
    out, ant = [], None
    for b in blocos:
        t = b["tipo"]
        if t == "h":
            s = "#" * b["nivel"] + " " + b["texto"]
        elif t == "li":
            s = "  " * b.get("nivel", 0) + "- " + b["texto"]
        elif t == "ol":
            s = "  " * b.get("nivel", 0) + f"{b.get('num', 1)}. " + b["texto"]
        elif t == "hr":
            s = "---"
        else:
            s = b["texto"]
        if not s.strip():
            continue
        junto = t in ("li", "ol") and ant in ("li", "ol")
        out.append(("\n" if junto else "\n\n") + s if out else s)
        ant = t
    return "".join(out).strip() + "\n"


# ============================================================
# TXT
# ============================================================
def decodificar(dados, reg):
    if dados.startswith(b"\xef\xbb\xbf"):
        dados = dados[3:]
        reg.add("marca BOM do início removida")
    try:
        return dados.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        pass
    try:
        from charset_normalizer import from_bytes
        m = from_bytes(dados).best()
        if m is not None:
            return str(m), m.encoding
    except ImportError:
        pass
    return dados.decode("cp1252", errors="replace"), "cp1252"


def parece_markdown(t):
    sinais = (len(re.findall(r"(?m)^#{1,6} \S", t)) * 3
              + len(re.findall(r"\*\*[^*\n]+\*\*", t))
              + len(re.findall(r"(?m)^\s*[-*] \S", t))
              + len(re.findall(r"\[[^\]\n]+\]\([^)\s]+\)", t)) * 2)
    return sinais >= 4


def converter_txt(orig, wd, reg):
    texto, enc = decodificar(orig.read_bytes(), reg)
    if enc.lower().replace("_", "-") not in ("utf-8", "ascii"):
        reg.add(f"codificação {enc} convertida para UTF-8")
    if "\r" in texto:
        texto = texto.replace("\r\n", "\n").replace("\r", "\n")
        reg.add("quebras de linha do Windows/Mac antigo normalizadas")
    texto = texto.replace("\f", "\n\n")
    if parece_markdown(texto):
        reg.avisos.append("o .txt já contém marcação Markdown (títulos, listas, negrito) — mantida como Markdown")
        return texto.rstrip() + "\n", {}

    segs = segmentos_com_carimbo(texto.strip(), reg)
    if segs is not None:  # transcrição copiada com carimbos de tempo: os segmentos cortam frases ao meio
        return juntar_segmentos(segs, reg), {}

    linhas = texto.strip().split("\n")  # linhas em branco no fim não contam como separador de parágrafos
    cheias = [l for l in linhas if l.strip()]
    tem_vazias = len(cheias) < len(linhas)
    if cheias and not tem_vazias and len(cheias) > 1:
        tams = sorted(len(l.rstrip()) for l in cheias)
        p75 = tams[int(len(tams) * .75)]
        quebrado = p75 < 110 and len(cheias) >= 8 and \
            sum(1 for t in tams if t >= p75 * .8) >= len(tams) * .5
        if quebrado:  # hard-wrap sem linhas em branco: parágrafo fecha em linha curta com pontuação final
            pars, atual = [], []
            for l in cheias:
                atual.append(l.strip())
                if re.search(r"[.!?:…»”\"]$", l.rstrip()) and len(l.rstrip()) < p75 * .8:
                    pars.append(" ".join(atual))
                    atual = []
            if atual:
                pars.append(" ".join(atual))
            reg.add("linhas quebradas a cada ~%d caracteres reunidas em parágrafos" % p75, n=len(pars))
        else:  # uma linha = um parágrafo
            pars = [l.strip() for l in cheias]
            reg.add("cada linha do .txt tratada como um parágrafo", n=len(pars))
        texto = "\n\n".join(pars)
    elif cheias:
        texto = re.sub(r"\n[ \t]*\n(?:[ \t]*\n)+", "\n\n", texto)

    texto = dehifenizar_linhas(texto, reg)
    return esc_texto(texto.strip()) + "\n", {}


def dehifenizar_linhas(texto, reg):
    """'pala-\\nvra' dentro de um parágrafo → 'palavra' (ou 'pala-vra' se o composto existe no texto)."""
    vocab = collections.Counter(w.lower() for w in re.findall(r"[^\W\d_]+(?:-[^\W\d_]+)*", texto))

    def troca(m):
        a, b = m.group(1), m.group(2)
        junto, hifen = a + b, f"{a}-{b}"
        if vocab.get(hifen.lower(), 0) > 0 and vocab.get(junto.lower(), 0) == 0:
            reg.add("hifenização de fim de linha desfeita (composto mantido)", f"{a}-⏎{b} → {hifen}")
            return hifen
        reg.add("hifenização de fim de linha desfeita", f"{a}-⏎{b} → {junto}")
        return junto
    return re.sub(r"([^\W\d_]+)-[ \t]*\n[ \t]*([a-zà-öø-ÿ][^\W\d_]*)", troca, texto)


# sem (?!\d) no fim: o rótulo acessível vem colado ('0:00' + '0 segundo')
_CARIMBO = re.compile(r"[ \t]*(?:(\d{1,2}):)?(\d{1,2}):(\d{2})")
_ROTULO_TEMPO = re.compile(
    r"(?:(\d+)\s*(?:hora|hour)(s?)[,\s]*(?:e\s+|and\s+|y\s+)?)?"
    r"(?:(\d+)\s*(?:minuto|minute)(s?)[,\s]*(?:e\s+|and\s+|y\s+)?)?"
    r"(?:(\d+)\s*(?:segundo|second)(s?))?", re.I)


def segmentos_com_carimbo(texto, reg):
    """Transcrição copiada do YouTube e afins: cada linha começa com um carimbo 'M:SS' (ou 'H:MM:SS'),
    às vezes colado ao rótulo acessível ('1:071 minuto e 7 segundos…'). Devolve os segmentos sem os
    carimbos, ou None se o texto não tem esse formato (≥ 60% das linhas com carimbo, mínimo de 5)."""
    linhas = [l for l in texto.split("\n") if l.strip()]
    if len(linhas) < 5 or sum(1 for l in linhas if _CARIMBO.match(l)) < len(linhas) * .6:
        return None
    segs, n, exemplos = [], 0, []
    for l in linhas:
        m = _CARIMBO.match(l)
        if not m:
            segs.append(l.strip())
            continue
        h, mi, s = int(m.group(1) or 0), int(m.group(2)), int(m.group(3))
        resto = l[m.end():]
        r = _ROTULO_TEMPO.match(resto)
        if r and r.group(0).strip():  # o rótulo só sai se os números batem com o carimbo
            g = r.groups()
            rh, rm, rs = (int(x) if x else 0 for x in g[0::2])
            if (rh, rm, rs) == (h, mi, s):
                fim = r.end()
                ultimo = max(i for i in (0, 2, 4) if g[i])
                if int(g[ultimo]) == 1 and g[ultimo + 1]:  # '1 segundo' + 'sabe': o 's' é do texto
                    fim -= 1
                resto = resto[fim:]
        n += 1
        if len(exemplos) < 2:
            exemplos.append(l[:len(l) - len(resto)].strip())
        if resto.strip():
            segs.append(resto.strip())
    reg.add("carimbos de tempo da transcrição removidos", " · ".join(exemplos), n=n)
    return segs


def juntar_segmentos(segs, reg):
    """Segmentos de transcrição (linhas com carimbo, blocos de legenda) viram texto corrido; anotações
    sonoras isoladas ([música], [aplausos]) ficam em parágrafo próprio."""
    pars, atual = [], []
    for s in segs:
        if re.fullmatch(r"\[[^\]]{1,40}\]", s):
            if atual:
                pars.append(" ".join(atual))
                atual = []
            pars.append(s)
        else:
            atual.append(s)
    if atual:
        pars.append(" ".join(atual))
    reg.add("segmentos da transcrição reunidos em texto corrido (anotações sonoras isoladas ficam em parágrafo próprio)",
            n=len(pars))
    texto = dehifenizar_linhas("\n\n".join(pars), reg)
    return esc_texto(texto.strip()) + "\n"


_TEMPO_LEGENDA = re.compile(r"^\s*(?:\d+:)?\d{1,2}:\d{2}[.,]\d{3}\s*-->")


def converter_legenda(orig, wd, reg):
    """.srt / .vtt: descarta índices, tempos, cabeçalho WEBVTT, blocos NOTE/STYLE/REGION e marcação
    de estilo; linhas repetidas em sequência (legenda 'rolante' do YouTube) entram uma vez só."""
    texto, enc = decodificar(orig.read_bytes(), reg)
    if enc.lower().replace("_", "-") not in ("utf-8", "ascii"):
        reg.add(f"codificação {enc} convertida para UTF-8")
    texto = texto.lstrip("﻿").replace("\r\n", "\n").replace("\r", "\n")
    segs, n_blocos, n_marc, n_rep = [], 0, 0, 0
    for bloco in re.split(r"\n\s*\n", texto):
        linhas = [l for l in bloco.split("\n") if l.strip()]
        if not linhas or re.match(r"(WEBVTT|NOTE|STYLE|REGION)\b", linhas[0]):
            continue
        k = next((i for i, l in enumerate(linhas) if _TEMPO_LEGENDA.match(l)), None)
        if k is None:
            continue
        n_blocos += 1
        for l in linhas[k + 1:]:
            limpa = re.sub(r"<[^>\n]*>|\{\\[^}\n]*\}", "", l)
            if limpa != l:
                n_marc += 1
            limpa = " ".join(html.unescape(limpa).split())
            if not limpa:
                continue
            if segs and segs[-1] == limpa:
                n_rep += 1
                continue
            segs.append(limpa)
    if not segs:
        raise ErroConversao("nenhum bloco de legenda encontrado (esperado: linhas 'início --> fim' seguidas do texto)")
    reg.add("índices e tempos das legendas removidos", n=n_blocos)
    if n_marc:
        reg.add("marcação de estilo das legendas removida (<i>, <c>, {\\an8}…)", n=n_marc)
    if n_rep:
        reg.add("linhas repetidas da legenda rolante removidas", n=n_rep)
    return juntar_segmentos(segs, reg), {}


# ============================================================
# DOCX
# ============================================================
def converter_docx(orig, wd, reg):
    try:
        import docx
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError:
        raise ErroConversao("python-docx não instalado — rode `gerar setup`")
    try:
        d = docx.Document(str(orig))
    except Exception as e:
        raise ErroConversao(f"não consegui abrir o .docx ({e}). Se for .doc antigo, salve como .docx no Word.")

    meta = {"titulo": (d.core_properties.title or "").strip(), "autor": (d.core_properties.author or "").strip(),
            "assunto": (d.core_properties.subject or "").strip()}
    figs = wd / "figuras"
    n_img = [0]

    # numeração: numId → ilvl → formato (bullet/decimal)
    formatos_num = {}
    try:
        numbering = d.part.numbering_part.element
        abstratos = {}
        for an in numbering.findall(qn("w:abstractNum")):
            niveis = {}
            for lvl in an.findall(qn("w:lvl")):
                fmt = lvl.find(qn("w:numFmt"))
                niveis[lvl.get(qn("w:ilvl"))] = fmt.get(qn("w:val")) if fmt is not None else "bullet"
            abstratos[an.get(qn("w:abstractNumId"))] = niveis
        for num in numbering.findall(qn("w:num")):
            ref = num.find(qn("w:abstractNumId"))
            if ref is not None:
                formatos_num[num.get(qn("w:numId"))] = abstratos.get(ref.get(qn("w:val")), {})
    except (NotImplementedError, AttributeError, KeyError):
        pass

    notas = {}
    try:
        for rel in d.part.rels.values():
            if rel.reltype.endswith("/footnotes"):
                from lxml import etree
                raiz = etree.fromstring(rel.target_part.blob)
                for fn in raiz.findall(qn("w:footnote")):
                    fid = fn.get(qn("w:id"))
                    txt = "".join(t.text or "" for t in fn.iter(qn("w:t"))).strip()
                    if txt and fid not in ("-1", "0"):
                        notas[fid] = txt
    except Exception:
        reg.avisos.append("notas de rodapé do .docx não puderam ser lidas")

    def estilo(p):
        return ((p.style.name if p.style is not None else "") or "")

    def nivel_titulo(p):
        nome = (p.style.name if p.style is not None else "") or ""
        if nome in ("Title", "Título"):
            return 1
        m = re.match(r"(?:Heading|Título|Titulo|Cabeçalho)\s*(\d)", nome, re.I)
        if m:
            return min(6, int(m.group(1)) + (1 if tem_title[0] else 0))
        ol = p._p.find(f"{qn('w:pPr')}/{qn('w:outlineLvl')}")
        if ol is not None and ol.get(qn("w:val")) not in (None, "9"):
            return min(6, int(ol.get(qn("w:val"))) + 1 + (1 if tem_title[0] else 0))
        return 0

    tem_title = [any((p.style is not None and p.style.name in ("Title", "Título")) for p in d.paragraphs)]

    def texto_run(r, fmt_ok=True):
        partes = []
        for el in r._r:
            tag = el.tag.split("}")[-1]
            if tag == "t":
                partes.append(("t", el.text or ""))
            elif tag in ("br", "cr"):
                partes.append(("br", ""))
            elif tag == "tab":
                partes.append(("t", " "))
            elif tag == "footnoteReference":
                partes.append(("fn", el.get(qn("w:id"))))
            elif tag == "drawing":
                for blip in el.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip"):
                    rid = blip.get(qn("r:embed"))
                    if rid and rid in d.part.related_parts:
                        parte = d.part.related_parts[rid]
                        n_img[0] += 1
                        ext = Path(parte.partname).suffix or ".png"
                        figs.mkdir(exist_ok=True)
                        nome = f"docx-{n_img[0]:02d}{ext}"
                        (figs / nome).write_bytes(parte.blob)
                        partes.append(("img", f"figuras/{nome}"))
        return partes

    def inline(p):
        """Runs → Markdown inline; negrito/itálico agrupados; imagens devolvidas à parte."""
        segs, imgs = [], []
        itens = p.iter_inner_content() if hasattr(p, "iter_inner_content") else p.runs
        for it in itens:
            if hasattr(it, "runs") and not hasattr(it, "_r"):  # Hyperlink
                url = getattr(it, "url", "") or ""
                txt = "".join(r.text for r in it.runs)
                if txt:
                    segs.append((f"[{esc_inline(txt)}]({url})" if url.startswith("http") else esc_inline(txt),
                                 False, False, True))
                continue
            b, i = bool(it.bold), bool(it.italic)
            for tipo, val in texto_run(it):
                if tipo == "t":
                    segs.append((val, b, i, False))
                elif tipo == "br":
                    segs.append(("\n", False, False, False))
                elif tipo == "fn" and val in notas:
                    segs.append((f"[^{val}]", False, False, True))
                elif tipo == "img":
                    imgs.append(val)
        # agrupa trechos contíguos com o mesmo estilo
        grupos = []
        for s, b, i, cru in segs:
            if grupos and grupos[-1][1:] == (b, i, cru) and not cru and s != "\n":
                grupos[-1][0] += s
            else:
                grupos.append([s, b, i, cru])
        out = []
        for s, b, i, cru in grupos:
            if cru:
                out.append(s)
                continue
            if s == "\n":
                out.append("\\\n")
                continue
            if not s.strip():
                out.append(s)
                continue
            lead = s[:len(s) - len(s.lstrip())]
            trail = s[len(s.rstrip()):]
            core = esc_inline(s.strip())
            if b and i:
                core = f"***{core}***"
            elif b:
                core = f"**{core}**"
            elif i:
                core = f"*{core}*"
            out.append(lead + core + trail)
        txt = "".join(out).strip()
        txt = re.sub(r"(\\\n)+$", "", txt).strip()
        return txt, imgs

    blocos = []
    for el in d.element.body.iterchildren():
        tag = el.tag.split("}")[-1]
        if tag == "p":
            p = Paragraph(el, d)
            txt, imgs = inline(p)
            nv = nivel_titulo(p)
            numpr = el.find(f"{qn('w:pPr')}/{qn('w:numPr')}")
            if txt:
                if nv:
                    blocos.append({"tipo": "h", "nivel": nv, "texto": re.sub(r"\*+|\\\n", " ", txt).strip()})
                elif numpr is not None:
                    ilvl_el = numpr.find(qn("w:ilvl"))
                    numid_el = numpr.find(qn("w:numId"))
                    ilvl = ilvl_el.get(qn("w:val")) if ilvl_el is not None else "0"
                    fmt = formatos_num.get(numid_el.get(qn("w:val")) if numid_el is not None else "", {}).get(ilvl, "bullet")
                    tipo = "li" if fmt in ("bullet", "none") else "ol"
                    blocos.append({"tipo": tipo, "nivel": int(ilvl), "texto": txt})
                elif re.search(r"List (Bullet|Number|Paragraph)|Lista|Marcador|Numera", estilo(p), re.I):
                    num = re.search(r"Number|Numera", estilo(p), re.I) is not None
                    m_nv = re.search(r"(\d)$", estilo(p))
                    blocos.append({"tipo": "ol" if num else "li",
                                   "nivel": int(m_nv.group(1)) - 1 if m_nv else 0, "texto": txt})
                else:
                    blocos.append({"tipo": "p", "texto": esc_bloco(txt)})
            for src in imgs:
                blocos.append({"tipo": "img", "texto": f"![]({src})"})
        elif tag == "tbl":
            t = Table(el, d)
            linhas = []
            for row in t.rows:
                cels = []
                for c in row.cells:
                    cels.append(" ".join(esc_inline(pp.text.strip()) for pp in c.paragraphs if pp.text.strip())
                                .replace("|", "\\|"))
                linhas.append(cels)
            if not linhas:
                continue
            n = max(len(l) for l in linhas)
            linhas = [l + [""] * (n - len(l)) for l in linhas]
            md = ["| " + " | ".join(linhas[0]) + " |", "|" + "---|" * n]
            md += ["| " + " | ".join(l) + " |" for l in linhas[1:]]
            blocos.append({"tipo": "tabela", "texto": "\n".join(md)})
            reg.add("tabela do .docx convertida em tabela Markdown")

    # numeração contínua das listas ordenadas
    cont = {}
    for b in blocos:
        if b["tipo"] == "ol":
            cont[b["nivel"]] = cont.get(b["nivel"], 0) + 1
            b["num"] = cont[b["nivel"]]
        elif b["tipo"] != "li":
            cont = {}
    md = montar_md(blocos)
    if notas:
        md += "\n" + "\n".join(f"[^{k}]: {esc_inline(v)}" for k, v in notas.items()) + "\n"
        reg.add("notas de rodapé do .docx convertidas em notas Markdown", n=len(notas))
    if n_img[0]:
        reg.add("imagens do .docx extraídas para figuras/", n=n_img[0])
    n_h = sum(1 for b in blocos if b["tipo"] == "h")
    if n_h:
        reg.add("estilos de título do Word convertidos em títulos Markdown", n=n_h)
    return md, meta


# ============================================================
# PDF
# ============================================================
RE_FOLIO = re.compile(r"^\s*(?:p[áa]g(?:ina)?\.?\s*)?(\d{1,4}|[ivxlcdm]{1,7})(?:\s*(?:de|/|of)\s*\d{1,4})?\s*$", re.I)
RE_SUMARIO = re.compile(r"(?:\.\s?){4,}\s*\d{1,4}\s*$|\s{3,}\d{1,4}\s*$")
RE_MARCADOR = re.compile(r"^\s*([•◦▪▫■□●○‣⁃–—-])\s+")
LIGADURAS = {"ﬀ": "ff", "ﬁ": "fi", "ﬂ": "fl", "ﬃ": "ffi", "ﬄ": "ffl", "ﬅ": "st", "ﬆ": "st"}


def _estilo(fontname):
    f = (fontname or "").split("+")[-1].lower()
    return ("bold" in f or "black" in f or "heavy" in f or "semibold" in f or "demi" in f,
            "italic" in f or "oblique" in f or re.search(r"[-,]it\b|-it$|italic", f) is not None)


def _chars_pdfium(pdf_path, idx):
    """Caracteres da página pelo pdfium, no formato do pdfplumber. Usado quando o pdfminer não
    consegue ler uma fonte e descarta o texto em silêncio."""
    import pypdfium2 as pdfium
    import pypdfium2.raw as raw
    doc = pdfium.PdfDocument(str(pdf_path))
    try:
        pg = doc[idx]
        alt = pg.get_height()
        tp = pg.get_textpage()
        out = []
        for i in range(tp.count_chars()):
            t = tp.get_text_range(i, 1)
            if not t or t in "\r\n":
                continue
            l, b, r, tt = tp.get_charbox(i)
            out.append({"text": t, "x0": l, "x1": r, "top": alt - tt, "bottom": alt - b,
                        "size": raw.FPDFText_GetFontSize(tp.raw, i) or (tt - b), "fontname": ""})
        return out
    finally:
        doc.close()


def _texto_pdfium(pdf_path, idx):
    import pypdfium2 as pdfium
    doc = pdfium.PdfDocument(str(pdf_path))
    try:
        return doc[idx].get_textpage().get_text_range()
    finally:
        doc.close()


def _linhas_da_pagina(page, chars=None):
    """Reconstrói linhas a partir dos caracteres: espaço onde o vão passa do espaçamento típico da linha
    (resolve texto com tracking, como 'S U M Á R I O')."""
    chars = [c for c in (chars if chars is not None else page.chars) if c.get("text") is not None]
    if not chars:
        return []
    # capitulares: letra isolada muito maior que o texto da página — viram prefixo da linha ao lado
    tam_pag = statistics.median(c["size"] for c in chars if c["text"].strip()) if any(
        c["text"].strip() for c in chars) else 10
    capitulares = [c for c in chars if len(c["text"]) == 1 and c["text"].isalpha() and c["size"] >= tam_pag * 2.2]
    if len(capitulares) <= 3:
        chars = [c for c in chars if c not in capitulares]
    else:
        capitulares = []
    # agrupa por sobreposição vertical do centro (fontes de reserva têm métricas diferentes:
    # o hífen de fim de linha, por exemplo, não pode escorregar para a linha seguinte)
    chars.sort(key=lambda c: (c["top"] + c["bottom"]) / 2)
    brutas = []
    for c in chars:
        cy = (c["top"] + c["bottom"]) / 2
        alvo = None
        for g in reversed(brutas[-4:]):
            if abs(g["cy"] - cy) <= max(1.5, min(g["tam"], c["size"]) * .45):
                alvo = g
                break
        if alvo is None:
            brutas.append({"cy": cy, "tam": c["size"], "chars": [c]})
        else:
            alvo["chars"].append(c)
    brutas = [g["chars"] for g in brutas]
    linhas = []
    for grupo in brutas:
        grupo.sort(key=lambda c: c["x0"])
        visiveis = [c for c in grupo if c["text"].strip()]
        if not visiveis:
            continue
        gaps = [b["x0"] - a["x1"] for a, b in zip(visiveis, visiveis[1:])]
        tam = statistics.median(c["size"] for c in visiveis)
        base = statistics.median(gaps) if gaps else 0
        base = max(0.0, min(base, tam * .6))
        limiar = base + tam * .17
        palavras, atual, ant = [], [], None
        for c in grupo:
            if not c["text"].strip():
                if atual:
                    palavras.append(atual)
                    atual = []
                ant = None
                continue
            if ant is not None and c["x0"] - ant["x1"] > limiar and atual:
                palavras.append(atual)
                atual = []
            atual.append(c)
            ant = c
        if atual:
            palavras.append(atual)
        ws = []
        for p in palavras:
            txt = "".join(c["text"] for c in p)
            bs = [_estilo(c.get("fontname")) for c in p]
            ws.append({"t": txt, "b": sum(b for b, _ in bs) > len(bs) / 2, "i": sum(i for _, i in bs) > len(bs) / 2})
        linhas.append({"top": min(c["top"] for c in visiveis), "bottom": max(c["bottom"] for c in visiveis),
                       "x0": visiveis[0]["x0"], "x1": visiveis[-1]["x1"], "size": tam, "palavras": ws,
                       "texto": " ".join(w["t"] for w in ws)})
    linhas.sort(key=lambda l: (l["top"], l["x0"]))
    for cap in capitulares:
        viz = [l for l in linhas if l["top"] >= cap["top"] - 3 and l["x0"] >= cap["x1"] - 2
               and l["top"] <= cap["bottom"]]
        if viz and viz[0]["palavras"]:
            l = viz[0]
            l["palavras"][0]["t"] = cap["text"] + l["palavras"][0]["t"]
            l["texto"] = " ".join(w["t"] for w in l["palavras"])
            l["capitular"] = True
    return linhas


def _ocr_pagina(pdf_path, idx, idioma):
    import pypdfium2 as pdfium
    doc = pdfium.PdfDocument(str(pdf_path))
    try:
        img = doc[idx].render(scale=300 / 72).to_pil()
    finally:
        pass
    with tempfile.TemporaryDirectory() as tmp:
        png = Path(tmp) / "p.png"
        img.save(png)
        r = subprocess.run(["tesseract", str(png), "stdout", "-l", idioma, "--psm", "3"],
                           capture_output=True, text=True, encoding="utf-8")
    doc.close()
    if r.returncode != 0:
        raise ErroConversao(f"tesseract falhou na página {idx + 1}: {r.stderr.strip()[:300]}")
    return r.stdout


def idioma_ocr(pedido=None):
    if not shutil.which("tesseract"):
        return None
    r = subprocess.run(["tesseract", "--list-langs"], capture_output=True, text=True)
    langs = set(r.stdout.split()[1:]) if r.returncode == 0 else set()
    if pedido:
        falta = [l for l in pedido.split("+") if l not in langs]
        if falta:
            raise ErroConversao(f"idioma de OCR não instalado no tesseract: {', '.join(falta)}")
        return pedido
    escolha = [l for l in ("por", "eng") if l in langs]
    return "+".join(escolha) or (sorted(langs)[0] if langs else "eng")


def converter_pdf(orig, wd, reg, idioma=None):
    try:
        import pdfplumber
    except ImportError:
        raise ErroConversao("pdfplumber não instalado — rode `gerar setup`")
    try:
        pdf = pdfplumber.open(str(orig))
    except Exception as e:
        raise ErroConversao(f"não consegui abrir o PDF ({e}). Se estiver protegido por senha, remova a proteção antes.")

    md_meta = pdf.metadata or {}
    meta = {"titulo": str(md_meta.get("Title", "") or "").strip(), "autor": str(md_meta.get("Author", "") or "").strip(),
            "assunto": str(md_meta.get("Subject", "") or "").strip()}
    paginas = []   # [(page_idx, linhas, imagens, altura, largura)]
    ocr_idx = []
    for k, page in enumerate(pdf.pages):
        page = page.dedupe_chars()
        n_chars = sum(1 for c in page.chars if c["text"].strip())
        n_pdfium = sum(1 for ch in _texto_pdfium(orig, k) if not ch.isspace())
        area = page.width * page.height
        grande = any((im["x1"] - im["x0"]) * (im["bottom"] - im["top"]) > area * .5 for im in page.images)
        if max(n_chars, n_pdfium) < 25 and (grande or page.images):
            ocr_idx.append(k)
            paginas.append((k, None, [], page.height, page.width))
            continue
        if n_pdfium > max(40, n_chars * 1.3):
            linhas = _linhas_da_pagina(page, _chars_pdfium(orig, k))
            reg.add("páginas lidas pelo pdfium (o leitor principal não decodificou uma das fontes)", f"p. {k + 1}")
        else:
            linhas = _linhas_da_pagina(page)
        # linhas duplicadas (efeito de sombra/contorno)
        vistos, unicas = set(), []
        for l in linhas:
            chave = (round(l["top"]), l["texto"])
            if chave in vistos:
                reg.add("linhas duplicadas pelo desenho da página removidas", l["texto"])
                continue
            vistos.add(chave)
            unicas.append(l)
        imgs = []
        for im in page.images:
            w, h = im["x1"] - im["x0"], im["bottom"] - im["top"]
            if w < 50 or h < 50:
                continue
            if w * h > area * .85:
                reg.add("imagem de fundo de página ignorada")
                continue
            imgs.append(im)
        paginas.append((k, unicas, imgs, page.height, page.width))

    if ocr_idx:
        lang = idioma_ocr(idioma)
        if lang is None:
            raise ErroConversao(
                f"{len(ocr_idx)} de {len(pdf.pages)} páginas são imagem (PDF escaneado) e precisam de OCR, "
                "mas o tesseract não está instalado. Rode `gerar verificar` para ver como instalar.")
        reg.add(f"páginas escaneadas lidas por OCR (tesseract, idioma {lang})", n=len(ocr_idx))
        reg.avisos.append(f"OCR em {len(ocr_idx)} página(s): o reconhecimento pode trocar letras — "
                          "o modo 'correção de língua' é recomendado")

    # tamanho do corpo
    tams = collections.Counter()
    for _, linhas, *_ in paginas:
        for l in linhas or []:
            tams[round(l["size"] * 2) / 2] += len(l["texto"])
    corpo = tams.most_common(1)[0][0] if tams else 11.0

    # cabeçalhos e rodapés corridos: zona de margem + (repetição, fólio ou letra miúda)
    def norm_cab(s):
        return re.sub(r"\d+", "#", " ".join(s.lower().split()))
    freq = collections.Counter()
    for _, linhas, _, alt, _ in paginas:
        for l in linhas or []:
            if l["top"] < alt * .09 or l["bottom"] > alt * .91:
                freq[norm_cab(l["texto"])] += 1
    n_pag = max(1, len(paginas))

    blocos_pag = []
    for k, linhas, imgs, alt, larg in paginas:
        if linhas is None:
            blocos_pag.append((k, None, []))
            continue
        mantidas = []
        n_toc = sum(1 for l in linhas if RE_SUMARIO.search(l["texto"]) and len(l["texto"]) < 160)
        pagina_sumario = n_toc >= 3 and n_toc >= len(linhas) * .4
        for l in linhas:
            if pagina_sumario:
                reg.add("página de sumário do original removida (o ebook gera um sumário novo)", l["texto"])
                continue
            margem = l["top"] < alt * .09 or l["bottom"] > alt * .91
            if margem and RE_FOLIO.match(l["texto"]):
                reg.add("números de página (fólios) removidos", l["texto"])
                continue
            if margem and (freq[norm_cab(l["texto"])] >= max(2, n_pag * .05) or l["size"] <= corpo * .82):
                reg.add("cabeçalhos e rodapés corridos removidos", l["texto"])
                continue
            if RE_SUMARIO.search(l["texto"]) and len(l["texto"]) < 160:
                reg.add("linhas do sumário original removidas (o ebook gera um sumário novo)", l["texto"])
                continue
            mantidas.append(l)
        blocos_pag.append((k, mantidas, imgs))

    # vocabulário para decidir hifenização
    vocab = collections.Counter()
    for _, linhas, _ in blocos_pag:
        for l in linhas or []:
            for w in l["palavras"][1:-1]:
                vocab[w["t"].lower().strip(".,;:!?…()\"“”'’")] += 1

    # níveis de título por tamanho relativo
    tams_tit = sorted({round(l["size"]) for _, ls, _ in blocos_pag for l in (ls or [])
                       if l["size"] >= corpo * 1.18 and len(l["texto"]) <= 140}, reverse=True)
    nivel_de = {t: min(3, i + 1) for i, t in enumerate(tams_tit)}
    if len(tams_tit) > 3:  # só os três maiores viram níveis distintos
        nivel_de = {t: (1 if i == 0 else 2 if i == 1 else 3) for i, t in enumerate(tams_tit)}

    def txt_linha(l):
        """Palavras → Markdown inline com negrito/itálico agrupados."""
        todas_b = all(w["b"] for w in l["palavras"])
        todas_i = all(w["i"] for w in l["palavras"])
        out, estilo_ant, buf = [], None, []

        def fecha():
            if not buf:
                return
            s = esc_inline(" ".join(buf))
            b, i = estilo_ant
            if b and i:
                s = f"***{s}***"
            elif b:
                s = f"**{s}**"
            elif i:
                s = f"*{s}*"
            out.append(s)
        for w in l["palavras"]:
            est = (w["b"] and not todas_b, w["i"] and not todas_i)
            if est != estilo_ant:
                fecha()
                buf, estilo_ant = [], est
            buf.append(w["t"])
        fecha()
        return " ".join(out), todas_b, todas_i

    def emenda(a, b):
        """Junta o fim de uma linha ao começo da seguinte, desfazendo a hifenização."""
        a = re.sub(r"[\u2010\u00ad]$", "-", a)
        ma = re.search(r"([^\W\d_]+)(-?)$", a)
        mb = re.match(r"^([^\W\d_]+)", b)
        if ma and mb and b[:1].islower():
            p1, hif, p2 = ma.group(1), ma.group(2), mb.group(1)
            junto = (p1 + p2).lower()
            if hif:
                if vocab.get(f"{p1}-{p2}".lower(), 0) and not vocab.get(junto, 0):
                    reg.add("hifenização de fim de linha desfeita (composto mantido)", f"{p1}-⏎{p2}")
                    return a + b
                reg.add("hifenização de fim de linha desfeita", f"{p1}-⏎{p2} → {p1 + p2}")
                return a[:-1] + b
            # hifenização sem hífen extraível (comum em PDF diagramado): só com evidência no vocabulário
            if vocab.get(junto, 0) >= 1 and (vocab.get(p1.lower(), 0) == 0 or vocab.get(p2.lower(), 0) == 0):
                reg.add("palavras partidas no fim da linha reunidas", f"{p1}⏎{p2} → {p1 + p2}")
                return a + b
        return a + " " + b

    blocos = []
    atual = None      # parágrafo em construção: {"texto", "ultima"}
    n_img = 0
    figs = wd / "figuras"

    def fecha_par():
        nonlocal atual
        if atual and atual["texto"].strip():
            blocos.append({"tipo": "p", "texto": esc_bloco(atual["texto"].strip())})
        atual = None

    gap_tipico = []
    for _, linhas, _ in blocos_pag:
        for a, b in zip(linhas or [], (linhas or [])[1:]):
            if abs(a["size"] - corpo) < 1 and abs(b["size"] - corpo) < 1:
                g = b["top"] - a["bottom"]
                if 0 <= g < corpo * 3:
                    gap_tipico.append(g)
    gap = statistics.median(gap_tipico) if gap_tipico else corpo * .4
    esquerda = collections.Counter(round(l["x0"]) for _, ls, _ in blocos_pag for l in (ls or [])
                                   if abs(l["size"] - corpo) < 1)
    margem_esq = esquerda.most_common(1)[0][0] if esquerda else 0
    direita = [l["x1"] for _, ls, _ in blocos_pag for l in (ls or []) if abs(l["size"] - corpo) < 1]
    margem_dir = statistics.quantiles(direita, n=10)[-1] if len(direita) >= 10 else (max(direita) if direita else 0)
    larg_col = max(1.0, margem_dir - margem_esq)

    for k, linhas, imgs in blocos_pag:
        if linhas is None:  # OCR
            fecha_par()
            txt = _ocr_pagina(orig, k, idioma_ocr(idioma))
            txt = "\n".join(l for l in txt.split("\n") if not RE_FOLIO.match(l) or not l.strip())
            txt = dehifenizar_linhas(txt, reg)
            vocab_ocr = collections.Counter(w.lower() for w in re.findall(r"[^\W\d_]+", txt))
            for par in re.split(r"\n\s*\n", txt):
                par = " ".join(par.split())
                if len(re.findall(r"[^\W\d_]", par)) < 3:
                    if par:
                        reg.add("fragmentos de ruído do OCR descartados", par)
                    continue
                mc = re.match(r"^([A-ZÀ-Ý]) ([a-zà-ÿ]+)\b", par)
                if mc and vocab_ocr.get((mc.group(1) + mc.group(2)).lower(), 0) \
                        and vocab_ocr.get(mc.group(2).lower(), 0) <= 1:  # 1 = o próprio fragmento
                    reg.add("capitulares separadas pelo OCR reunidas", f"{mc.group(0)} → {mc.group(1)}{mc.group(2)}")
                    par = mc.group(1) + par[2:]
                if par:
                    m = RE_MARCADOR.match(par)
                    if m:
                        blocos.append({"tipo": "li", "texto": esc_inline(par[m.end():])})
                    else:
                        blocos.append({"tipo": "p", "texto": esc_bloco(esc_inline(par))})
            continue
        eventos = [("l", l["top"], l) for l in linhas] + [("i", im["top"], im) for im in imgs]
        eventos.sort(key=lambda e: e[1])
        ant = None
        for tipo, _, obj in eventos:
            if tipo == "i":
                fecha_par()
                n_img += 1
                figs.mkdir(exist_ok=True)
                nome = f"pdf-p{k + 1:03d}-{n_img:02d}.png"
                try:
                    pg = pdf.pages[k]
                    bbox = (max(0, obj["x0"]), max(0, obj["top"]), min(pg.width, obj["x1"]), min(pg.height, obj["bottom"]))
                    pg.crop(bbox).to_image(resolution=200).save(str(figs / nome))
                    blocos.append({"tipo": "img", "texto": f"![](figuras/{nome})"})
                except Exception:
                    reg.avisos.append(f"imagem da página {k + 1} não pôde ser extraída")
                ant = None
                continue
            l = obj
            texto, tb, _ = txt_linha(l)
            titulo = l["size"] >= corpo * 1.18 and len(l["texto"]) <= 140
            negrito_curto = tb and len(l["texto"]) <= 90 and not re.search(r"[.,;:]$", l["texto"]) \
                and abs(l["size"] - corpo) < 1.5
            if titulo:
                fecha_par()
                nv = nivel_de.get(round(l["size"]), 3)
                limpo = " ".join(w["t"] for w in l["palavras"])
                if blocos and blocos[-1]["tipo"] == "h" and blocos[-1]["nivel"] == nv and ant is not None \
                        and ant.get("titulo") and l["top"] - ant["bottom"] < l["size"]:
                    blocos[-1]["texto"] += " " + esc_inline(limpo)
                else:
                    blocos.append({"tipo": "h", "nivel": nv, "texto": esc_inline(limpo)})
                ant = {**l, "titulo": True}
                continue
            m = RE_MARCADOR.match(l["texto"])
            novo = atual is None or ant is None or ant.get("titulo") or l.get("capitular")
            if not novo:
                vao = l["top"] - ant["bottom"]
                curta = ant["x1"] < margem_esq + larg_col * .85
                recuo = l["x0"] > margem_esq + corpo * .8 and ant["x0"] <= margem_esq + 3
                fim_frase = re.search(r"[.!?:…»”\"]$", ant["texto"]) is not None
                if vao > gap * 1.8 + 2 or recuo or (curta and fim_frase) or m or negrito_curto \
                        or ant.get("negrito_curto"):
                    novo = True
            if novo:
                fecha_par()
            if m:
                atual = None
                blocos.append({"tipo": "li", "texto": esc_inline(l["texto"][m.end():])})
                ant = {**l}
                continue
            if negrito_curto and (atual is None):
                blocos.append({"tipo": "h", "nivel": 3 if tams_tit else 2, "texto": esc_inline(l["texto"])})
                reg.add("linhas curtas em negrito tratadas como intertítulos", l["texto"])
                ant = {**l, "titulo": True}
                continue
            if atual is None:
                atual = {"texto": texto}
            else:
                atual["texto"] = emenda(atual["texto"], texto)
            ant = {**l}
        # parágrafo continua na página seguinte só se a última linha não fechou frase curta
        if ant is not None and atual is not None:
            curta = ant["x1"] < margem_esq + larg_col * .85
            if curta and re.search(r"[.!?:…»”\"]$", ant["texto"]):
                fecha_par()
    fecha_par()
    pdf.close()

    # palavra hifenizada que atravessou a quebra de página/coluna: reúne os dois parágrafos
    # e parágrafo partido no meio da frase (quebra de página, caixa, OCR): termina sem pontuação e o
    # seguinte começa em minúscula
    fundidos = []
    for bl in blocos:
        if fundidos and bl["tipo"] == "p" and fundidos[-1]["tipo"] == "p" and bl["texto"][:1].islower():
            ant_txt = fundidos[-1]["texto"]
            if re.search(r"[^\W\d_][-\u2010\u00ad]$", ant_txt):
                fundidos[-1]["texto"] = emenda(ant_txt, bl["texto"])
                continue
            if not re.search(r"[.!?:;…»”\")\]*]$", ant_txt):
                fundidos[-1]["texto"] = ant_txt + " " + bl["texto"]
                reg.add("parágrafos partidos no meio da frase reunidos", f"…{ant_txt[-30:]} ⏎ {bl['texto'][:30]}…")
                continue
        fundidos.append(bl)
    blocos = fundidos

    md = montar_md(blocos)
    for lig, sub in LIGADURAS.items():
        if lig in md:
            reg.add("ligaduras tipográficas (ﬁ, ﬂ…) convertidas em letras", n=md.count(lig))
            md = md.replace(lig, sub)
    if "­" in md:
        reg.add("hifens invisíveis (soft hyphen) removidos", n=md.count("­"))
        md = md.replace("­", "")
    if n_img:
        reg.add("imagens do PDF extraídas para figuras/", n=n_img)
    n_h = sum(1 for b in blocos if b["tipo"] == "h")
    if n_h:
        reg.add("títulos reconhecidos pelo tamanho da fonte", n=n_h)
    if len(paginas) and not md.strip():
        raise ErroConversao("não foi possível extrair texto deste PDF")
    return md, meta


# ============================================================
# CLASSIFICAÇÃO DO TIPO DE FONTE
# ============================================================
ORALIDADE = re.compile(r"\b(né|tá|beleza|pessoal|a gente|tipo assim|vocês|gente|entendeu|olha só|"
                       r"vamos lá|deixa eu|tô|pra|daí|aí)\b", re.I)


def classificar(texto):
    palavras = len(re.findall(r"\w+", texto)) or 1
    titulos = re.findall(r"(?m)^(#{1,6}) \S", texto)
    niveis = {len(t) for t in titulos}
    blocos = [b for b in re.split(r"\n\s*\n", texto) if b.strip()]
    listas = sum(1 for b in blocos if re.match(r"\s*([-*+]|\d+[.)])\s", b))
    media_par = sum(len(b) for b in blocos) / max(1, len(blocos))
    oral = len(ORALIDADE.findall(texto)) * 1000 / palavras
    indicios = [f"{palavras:,} palavras".replace(",", "."), f"{len(titulos)} títulos em {len(niveis)} nível(is)",
                f"{oral:.1f} marcas de oralidade por mil palavras",
                f"parágrafo médio de {media_par:.0f} caracteres", f"{listas} blocos de lista"]
    if oral >= 25 or (oral >= 12 and len(titulos) < max(3, palavras / 4000)):
        tipo = "transcricao"
    elif len(titulos) >= 8 and (palavras > 12000 or len(niveis) >= 2):
        tipo = "livro"
    elif media_par < 160 or listas >= max(4, len(blocos) * .3):
        tipo = "notas"
    else:
        tipo = "documento"
    return tipo, indicios


TIPOS_DESCRICAO = {
    "transcricao": "transcrição de fala (aula, palestra, entrevista, podcast)",
    "documento": "documento escrito (artigo, relatório, ensaio, apostila curta)",
    "livro": "livro ou apostila longa com hierarquia de títulos",
    "notas": "notas, listas e tópicos curtos",
}


# ============================================================
# ENTRADA
# ============================================================
def converter(orig, wd, idioma=None):
    """Devolve (markdown, info). Nunca escreve no original."""
    fmt = FORMATOS.get(orig.suffix.lower())
    if fmt is None:
        raise ErroConversao(f"formato não suportado: {orig.suffix or '(sem extensão)'}. "
                            "Aceitos: .md, .txt, .docx, .pdf, .srt, .vtt. Para .doc/.rtf/.odt/.pages, salve como .docx.")
    reg = Registro()
    meta = {}
    if fmt == "markdown":
        dados = orig.read_bytes()
        try:
            md = dados.decode("utf-8")
            copia = True
        except UnicodeDecodeError:
            md, enc = decodificar(dados, reg)
            reg.add(f"codificação {enc} convertida para UTF-8")
            copia = False
        info_extra = {"copia_exata": copia}
    elif fmt == "texto":
        md, meta = converter_txt(orig, wd, reg)
        info_extra = {}
    elif fmt == "legenda":
        md, meta = converter_legenda(orig, wd, reg)
        info_extra = {}
    elif fmt == "docx":
        md, meta = converter_docx(orig, wd, reg)
        info_extra = {}
    else:
        md, meta = converter_pdf(orig, wd, reg, idioma)
        info_extra = {}
    md = unicodedata.normalize("NFC", md) if fmt != "markdown" else md
    tipo, indicios = classificar(md)
    if fmt == "legenda":  # legenda é fala por definição, mesmo com pouca marca de oralidade
        tipo = "transcricao"
    return md, {"formato": fmt, "metadados": {k: v for k, v in meta.items() if v},
                "limpeza": reg.lista(), "avisos": reg.avisos, "tipo_fonte": tipo,
                "tipo_descricao": TIPOS_DESCRICAO[tipo], "indicios": indicios, **info_extra}
